#!/usr/bin/env python3
"""Apply a restrained, journal-ready Word style to Online Resource 1."""

from pathlib import Path
import sys

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


NAVY = "1F4E78"
PALE_BLUE = "EAF2F8"
GRID = "D9D9D9"


def set_run_font(run, name="Times New Roman", size=None, bold=None, color=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def borders(cell, color=GRID):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_borders = tc_pr.first_child_found_in("w:tcBorders")
    if tc_borders is None:
        tc_borders = OxmlElement("w:tcBorders")
        tc_pr.append(tc_borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = qn(f"w:{edge}")
        element = tc_borders.find(tag)
        if element is None:
            element = OxmlElement(f"w:{edge}")
            tc_borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "4")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_table_layout(table, widths):
    tbl_pr = table._tbl.tblPr
    layout = tbl_pr.first_child_found_in("w:tblLayout")
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")
    for row in table.rows:
        for idx, width in enumerate(widths):
            cell = row.cells[idx]
            cell.width = Inches(width)
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(int(width * 1440)))
            tc_w.set(qn("w:type"), "dxa")


def keep_with_next(paragraph):
    p_pr = paragraph._p.get_or_add_pPr()
    p_pr.append(OxmlElement("w:keepNext"))


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = "PAGE"
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char1, instr_text, fld_char2])
    set_run_font(run, size=9, color="4F4F4F")


def main(source, output):
    doc = Document(source)
    section = doc.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)

    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.08

    for name, size in (("Title", 16), ("Heading 1", 13), ("Heading 2", 11)):
        style = doc.styles[name]
        style.font.name = "Times New Roman"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)

    doc.styles["Title"].paragraph_format.space_after = Pt(16)
    doc.styles["Heading 1"].paragraph_format.space_before = Pt(12)
    doc.styles["Heading 1"].paragraph_format.space_after = Pt(6)

    for p in doc.paragraphs:
        for run in p.runs:
            set_run_font(run)
        if p.style.name == "Title":
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if p.style.name.startswith("Heading"):
            keep_with_next(p)
        text = p.text.strip()
        if text.startswith(("Table S", "Fig. S")):
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(6)
            keep_with_next(p)
        if text == "Contents":
            p.paragraph_format.space_before = Pt(10)

    table_widths = (
        [1.40, 0.65, 1.15, 1.00, 0.40, 0.70, 0.70, 1.00],
        [1.10, 0.85, 0.85, 0.85, 1.00, 1.00, 1.25],
    )
    for table_index, table in enumerate(doc.tables):
        table.autofit = False
        if table_index < len(table_widths):
            set_table_layout(table, table_widths[table_index])
        if table_index == 0:
            table.rows[0].cells[1].text = "T"
        repeat_table_header(table.rows[0])
        for r_idx, row in enumerate(table.rows):
            for cell in row.cells:
                cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
                borders(cell)
                if r_idx == 0:
                    shade(cell, NAVY)
                elif r_idx % 2 == 0:
                    shade(cell, PALE_BLUE)
                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.space_after = Pt(2)
                    paragraph.paragraph_format.space_before = Pt(2)
                    for run in paragraph.runs:
                        set_run_font(run, size=9.2, bold=(r_idx == 0), color=("FFFFFF" if r_idx == 0 else None))

    header = section.header.paragraphs[0]
    header.text = "Online Resource 1 | Acropora palmata demography"
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for run in header.runs:
        set_run_font(run, size=8.5, color="4F4F4F")

    footer = section.footer.paragraphs[0]
    footer.clear()
    add_page_number(footer)

    for image in doc.inline_shapes:
        image._inline.docPr.set("descr", Path(getattr(image, "_inline").docPr.get("name", "Supplementary figure")).stem.replace("_", " "))

    doc.core_properties.title = "Online Resource 1 Supporting analyses for Acropora palmata demography"
    doc.core_properties.author = "Adrian C. Stier"
    doc.core_properties.subject = "Coral Reefs electronic supplementary material"
    doc.save(output)


if __name__ == "__main__":
    if len(sys.argv) != 3:
        raise SystemExit("Usage: format_supplement_docx.py input.docx output.docx")
    main(sys.argv[1], sys.argv[2])
